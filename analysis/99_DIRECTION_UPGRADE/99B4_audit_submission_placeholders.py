#!/usr/bin/env python3
"""Read-only structural text/placeholder audit of the current submission DOCX drafts."""

from __future__ import annotations

import argparse
import csv
import re
from pathlib import Path

from docx import Document


PATTERNS = {
    "explicit_placeholder": re.compile(r"\[(?:TO BE COMPLETED|INSERT[^\]]*|PLACEHOLDER|TBD|TODO)[^\]]*\]", re.I),
    "blank_field": re.compile(r"(?:^|\s)(?:_{3,}|\.\s*\.\s*\.)"),
    "author_information": re.compile(r"authors?|affiliations?|corresponding|ORCID|email|telephone|address", re.I),
    "compliance_information": re.compile(r"funding|conflict|competing interest|ethics?|consent|CRediT|acknowledg|preprint|generative AI|artificial intelligence", re.I),
    "venue_information": re.compile(r"editor|journal|manuscript type|word count|suggested reviewer|opposed reviewer", re.I),
    "availability_information": re.compile(r"data availability|code availability|Zenodo|DOI|GitHub|license|CITATION", re.I),
}


def blocks(document: Document):
    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = " ".join(paragraph.text.split())
        if text:
            yield f"P{index}", text
    for table_index, table in enumerate(document.tables, start=1):
        for row_index, row in enumerate(table.rows, start=1):
            values = [" ".join(cell.text.split()) for cell in row.cells]
            text = " | ".join(values)
            if text.strip(" |"):
                yield f"T{table_index}R{row_index}", text


def write_csv(path: Path, rows: list[dict[str, str]]) -> None:
    if path.exists():
        raise RuntimeError(f"Refusing to overwrite: {path}")
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--submission", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()
    submission = Path(args.submission).resolve()
    output = Path(args.output).resolve()
    output.mkdir(parents=True, exist_ok=True)

    extract_path = output / "submission_docx_structural_text_extract.txt"
    audit_path = output / "submission_docx_placeholder_audit.csv"
    if extract_path.exists() or audit_path.exists():
        raise RuntimeError("Refusing to overwrite an existing DOCX audit")

    documents = sorted((submission / "01_Manuscript").glob("*.docx"))
    if not documents:
        raise RuntimeError("No manuscript DOCX drafts found")
    extracted: list[str] = []
    audit_rows: list[dict[str, str]] = []
    for path in documents:
        document = Document(path)
        extracted.extend([f"=== {path.name} ===", ""])
        for location, text in blocks(document):
            extracted.append(f"[{location}] {text}")
            matched = [label for label, pattern in PATTERNS.items() if pattern.search(text)]
            if matched:
                audit_rows.append(
                    {
                        "document": path.name,
                        "location": location,
                        "flag_types": "|".join(matched),
                        "text": text,
                    }
                )
        extracted.append("")

    extract_path.write_text("\n".join(extracted), encoding="utf-8")
    if not audit_rows:
        audit_rows.append({"document": "", "location": "", "flag_types": "NONE", "text": "No candidate flags detected"})
    write_csv(audit_path, audit_rows)
    print(f"Audited {len(documents)} DOCX drafts; candidate rows={len(audit_rows)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
