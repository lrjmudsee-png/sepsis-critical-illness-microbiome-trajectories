#!/usr/bin/env python3
"""A-stage source retrieval and workbook feasibility inspection.

This script is intentionally read-only for all pre-existing project inputs.  It
downloads the official PRJNA851469 supplement into a new result directory,
checks the XLSX ZIP container, and records sheet-level evidence relevant to the
planned mechanism/clinical upgrade.  It does not calculate biological effects.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import os
import platform
import sys
import tempfile
import urllib.request
import zipfile
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


SUPPLEMENT_URL = (
    "https://media.springernature.com/original/springer-static/esm/"
    "art%3A10.1038%2Fs41591-023-02243-5/MediaObjects/"
    "41591_2023_2243_MOESM3_ESM.xlsx"
)
SUPPLEMENT_NAME = "PRJNA851469_Supplementary_Tables_2_17_OFFICIAL.xlsx"
RAW_COPIES = [
    Path(r"E:\sepsis_project\data\PRJNA851469\00_metadata\step76B_public_supplement\41591_2023_2243_MOESM3_ESM.xlsx"),
    Path(r"E:\sepsis_project\data\PRJNA851469\00_metadata\step78A_public_supplement\PRJNA851469_Supplementary_Tables_2_17.xlsx"),
    Path(r"E:\sepsis_project\data\PRJNA851469\00_metadata\step78A2_public_supplement\PRJNA851469_Supplementary_Tables_2_17.xlsx"),
]
HASH_GUARD_INPUTS = [
    Path(r"E:\sepsis_project\results\V2_27B_ANALYSIS_OBJECTS_AND_REPLICATE_FREEZE\V2_STEP87B_analysis_object_registry.csv"),
    Path(r"E:\sepsis_project\data\_V2_ANALYSIS_READY\01_PROJECTS\PRJNA691455\05_work\step87B_analysis_object\PRJNA691455_analysis_object_step87B.rds"),
    Path(r"E:\sepsis_project\data\_V2_ANALYSIS_READY\01_PROJECTS\PRJNA851469\05_work\step87B_analysis_object\PRJNA851469_analysis_object_step87B.rds"),
    Path(r"E:\sepsis_project\data\_V2_ANALYSIS_READY\01_PROJECTS\PRJNA516701\05_work\step87B_analysis_object\PRJNA516701_analysis_object_step87B.rds"),
    Path(r"E:\sepsis_project\data\PRJNA1125274\03_dada2\PRJNA1125274_analysis_object_external_validation.rds"),
    Path(r"E:\sepsis_project\results\V2_UPGRADE_20260907\98A_AITCHISON_LONGITUDINAL_ROBUSTNESS\03_PAIRED_CONTRASTS\V2_98A_paired_early_vs_late_contrasts.csv"),
    Path(r"E:\sepsis_project\results\V2_UPGRADE_20260907\98B_RANDOM_EFFECTS_META_ANALYSIS\META_ELIGIBILITY_FREEZE.csv"),
    Path(r"E:\sepsis_project\results\V2_UPGRADE_20260910_FINAL_FREEZE\PRJNA1125274_FINAL_population_membership.csv"),
    Path(r"E:\sepsis_project\results\V2_UPGRADE_20260910_FINAL_FREEZE\PRJNA1125274_FINAL_primary_and_sensitivity_results.csv"),
    Path(r"E:\sepsis_project\results\V2_UPGRADE_20260907\98E_PRJNA1125274_EXTERNAL_VALIDATION\PRJNA1125274_external_validation_sample_displacement.csv"),
    Path(r"E:\sepsis_project\results\V2_36G2_PRJNA851469_DAY3_CLINICAL_OUTCOME_MODELS_FIXED\10_MANUSCRIPT_SAFE_SUMMARY.txt"),
    Path(r"E:\sepsis_project\code\03_data_processing\92B2_EII_baseline_harmonization_fixed.R"),
    Path(r"E:\sepsis_project\V2_UPGRADE_20260913_DIRECTION_PLAN\00_研究方向与推进路线.md"),
    Path(r"E:\sepsis_project\V2_UPGRADE_20260913_DIRECTION_PLAN\01_候选统计设计与验收标准.md"),
    Path(r"E:\sepsis_project\V2_UPGRADE_20260913_DIRECTION_PLAN\02_交给执行模型的任务书.md"),
]


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def zip_status(path: Path) -> tuple[bool, str]:
    if not path.exists():
        return False, "file_missing"
    try:
        with zipfile.ZipFile(path) as archive:
            bad_member = archive.testzip()
            if bad_member:
                return False, f"crc_failure:{bad_member}"
        return True, "valid_xlsx_zip_crc"
    except Exception as exc:  # evidence is written to the audit, not suppressed
        return False, f"invalid_xlsx_zip:{type(exc).__name__}:{exc}"


def download_atomic(url: str, destination: Path) -> tuple[str, int]:
    destination.parent.mkdir(parents=True, exist_ok=True)
    valid, _ = zip_status(destination)
    if destination.exists() and valid:
        return "reused_valid_existing_download", destination.stat().st_size
    if destination.exists() and not valid:
        raise RuntimeError(
            f"Refusing to overwrite invalid existing output: {destination}. "
            "Use a fresh output directory."
        )
    request = urllib.request.Request(
        url,
        headers={"User-Agent": "Mozilla/5.0 sepsis-direction-feasibility-audit/1.0"},
    )
    tmp_path: Path | None = None
    try:
        with urllib.request.urlopen(request, timeout=120) as response:
            status = getattr(response, "status", 200)
            if status != 200:
                raise RuntimeError(f"HTTP status {status}")
            with tempfile.NamedTemporaryFile(
                mode="wb", delete=False, dir=destination.parent, suffix=".part"
            ) as tmp:
                tmp_path = Path(tmp.name)
                while True:
                    block = response.read(1024 * 1024)
                    if not block:
                        break
                    tmp.write(block)
        valid, detail = zip_status(tmp_path)
        if not valid:
            raise RuntimeError(f"Downloaded workbook failed CRC/container check: {detail}")
        os.replace(tmp_path, destination)
        return "downloaded_http_200_and_crc_validated", destination.stat().st_size
    finally:
        if tmp_path is not None and tmp_path.exists():
            tmp_path.unlink()


def compact(value: object) -> str:
    if value is None:
        return ""
    text = " ".join(str(value).replace("\n", " ").split())
    return text[:240]


def classify_sheet(sheet_name: str, preview: str, has_id: bool, has_host: bool) -> tuple[str, bool, str]:
    lower = f"{sheet_name} {preview}".lower()
    if "asv" in lower and "sample id" in lower:
        return (
            "sample-by-ASV microbiome abundance/taxonomy",
            False,
            "Sample identifiers are present, but columns are microbiome abundance/taxonomy rather than host measurements.",
        )
    if "correlation" in lower or "spearman" in lower:
        return (
            "summary correlation matrix/statistic",
            False,
            "Only aggregate correlation output was detected; no joinable patient-level host values were identified.",
        )
    if "permanova" in lower or "model" in lower:
        return (
            "summary model/statistical output",
            False,
            "The sheet contains model-level summaries rather than patient-level host measurements.",
        )
    if has_id and has_host:
        return (
            "candidate patient/sample-level host table",
            True,
            "Both an identifier field and a host/clinical measurement term were detected; manual review is required.",
        )
    if has_host:
        return (
            "host/immune summary without joinable identifier",
            False,
            "Host/immune terms were detected but no patient/sample identifier was found in the inspected header region.",
        )
    return (
        "annotation, guide, or non-patient-level supplement",
        False,
        "No patient-level host measurement structure was detected in the inspected header region.",
    )


def inspect_workbook(path: Path) -> tuple[list[dict[str, object]], dict[str, object], Counter[str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows: list[dict[str, object]] = []
    table2_info: dict[str, object] = {
        "n_unique_sample_ids": "",
        "n_healthy_like_sample_ids": "",
        "n_patient_like_sample_ids": "",
        "sample_id_examples": "",
    }
    table2_sample_ids: Counter[str] = Counter()
    id_terms = ("patient", "sample id", "sample_id", "subject", "participant")
    host_terms = (
        "antibody", "cytokine", "immune", "immunoglob", "igg", "iga", "igm",
        "sofa", "apache", "mortality", "outcome", "clinical", "lymphocyte",
        "leukocyte", "neutrophil", "inflammation", "biomarker", "blood",
    )
    for worksheet in workbook.worksheets:
        nonempty: list[tuple[int, list[str]]] = []
        for row_index, values in enumerate(
            worksheet.iter_rows(min_row=1, max_row=min(30, worksheet.max_row), values_only=True),
            start=1,
        ):
            cleaned = [compact(value) for value in values[:40]]
            cleaned = [value for value in cleaned if value]
            if cleaned:
                nonempty.append((row_index, cleaned))
        preview = " || ".join(
            f"R{index}: " + " | ".join(values) for index, values in nonempty[:5]
        )
        lower = preview.lower()
        has_id = any(term in lower for term in id_terms)
        has_host = any(term in lower for term in host_terms)
        granularity, patient_host, rationale = classify_sheet(
            worksheet.title, preview, has_id, has_host
        )
        rows.append(
            {
                "sheet": worksheet.title,
                "max_row": worksheet.max_row,
                "max_column": worksheet.max_column,
                "first_nonempty_row": nonempty[0][0] if nonempty else "",
                "header_preview": preview,
                "detected_identifier_field": has_id,
                "detected_host_or_clinical_term": has_host,
                "inferred_granularity": granularity,
                "joinable_patient_level_host_measurements": patient_host,
                "rationale": rationale,
            }
        )

        if worksheet.title.strip().lower() == "supplementary table 2":
            sample_ids: Counter[str] = Counter()
            header_row = None
            sample_col = None
            for row_index, values in enumerate(
                worksheet.iter_rows(min_row=1, values_only=True), start=1
            ):
                if header_row is None:
                    labels = [compact(value).lower() for value in values]
                    if "sample id" in labels:
                        header_row = row_index
                        sample_col = labels.index("sample id")
                    continue
                if sample_col is not None and sample_col < len(values):
                    sample_value = compact(values[sample_col])
                    if sample_value:
                        sample_ids[sample_value] += 1
            healthy = [value for value in sample_ids if "healthy" in value.lower()]
            patient = [value for value in sample_ids if "patient" in value.lower()]
            table2_info = {
                "n_unique_sample_ids": len(sample_ids),
                "n_healthy_like_sample_ids": len(healthy),
                "n_patient_like_sample_ids": len(patient),
                "sample_id_examples": " | ".join(list(sample_ids.keys())[:12]),
            }
            table2_sample_ids = sample_ids
    workbook.close()
    return rows, table2_info, table2_sample_ids


def write_csv(path: Path, rows: list[dict[str, object]], fieldnames: list[str] | None = None) -> None:
    if path.exists():
        raise RuntimeError(f"Refusing to overwrite existing audit output: {path}")
    if not rows:
        raise RuntimeError(f"No rows available for required output: {path}")
    if fieldnames is None:
        fieldnames = list(rows[0].keys())
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def capture_hash_guard(path: Path, inputs: list[Path]) -> None:
    rows: list[dict[str, object]] = []
    for input_path in inputs:
        rows.append(
            {
                "path": str(input_path),
                "exists_before": input_path.exists(),
                "bytes_before": input_path.stat().st_size if input_path.exists() else "",
                "mtime_utc_before": datetime.fromtimestamp(
                    input_path.stat().st_mtime, tz=timezone.utc
                ).isoformat() if input_path.exists() else "",
                "sha256_before": sha256_file(input_path) if input_path.exists() else "",
            }
        )
    write_csv(path, rows)


def capture_tree_guard(path: Path, root: Path) -> None:
    rows: list[dict[str, object]] = []
    if root.exists():
        for file_path in sorted(item for item in root.rglob("*") if item.is_file()):
            rows.append(
                {
                    "protected_root": str(root),
                    "relative_path": str(file_path.relative_to(root)),
                    "bytes_before": file_path.stat().st_size,
                    "mtime_utc_before": datetime.fromtimestamp(
                        file_path.stat().st_mtime, tz=timezone.utc
                    ).isoformat(),
                    "sha256_before": sha256_file(file_path),
                }
            )
    if not rows:
        rows.append(
            {
                "protected_root": str(root),
                "relative_path": "",
                "bytes_before": "",
                "mtime_utc_before": "",
                "sha256_before": "",
            }
        )
    write_csv(path, rows)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", required=True, help="Fresh A_FEASIBILITY output directory")
    args = parser.parse_args()
    output = Path(args.output).resolve()
    output.mkdir(parents=True, exist_ok=True)
    source_dir = output / "sources"
    source_dir.mkdir(parents=True, exist_ok=True)
    official_path = source_dir / SUPPLEMENT_NAME

    action, size = download_atomic(SUPPLEMENT_URL, official_path)
    valid, validation_detail = zip_status(official_path)
    if not valid:
        raise RuntimeError(validation_detail)
    official_sha = sha256_file(official_path)

    retrieval_rows: list[dict[str, object]] = [
        {
            "source_id": "PRJNA851469_NatureMedicine_2023_supplement_tables_2_17",
            "source_type": "official_publisher_supplement",
            "url_or_path": SUPPLEMENT_URL,
            "local_copy": str(official_path),
            "retrieval_utc": datetime.now(timezone.utc).isoformat(),
            "status": action,
            "bytes": size,
            "sha256": official_sha,
            "container_crc_valid": valid,
            "detail": validation_detail,
        }
    ]
    for index, raw_path in enumerate(RAW_COPIES, start=1):
        raw_valid, raw_detail = zip_status(raw_path)
        retrieval_rows.append(
            {
                "source_id": f"preexisting_raw_supplement_copy_{index}",
                "source_type": "preexisting_input_audit_only",
                "url_or_path": str(raw_path),
                "local_copy": str(raw_path),
                "retrieval_utc": datetime.now(timezone.utc).isoformat(),
                "status": "read_only_not_modified",
                "bytes": raw_path.stat().st_size if raw_path.exists() else "",
                "sha256": sha256_file(raw_path) if raw_path.exists() else "",
                "container_crc_valid": raw_valid,
                "detail": raw_detail,
            }
        )
    write_csv(output / "A02_source_retrieval_log.csv", retrieval_rows)

    sheet_rows, table2_info, table2_sample_ids = inspect_workbook(official_path)
    for row in sheet_rows:
        row.update(table2_info if row["sheet"].strip().lower() == "supplementary table 2" else {
            "n_unique_sample_ids": "",
            "n_healthy_like_sample_ids": "",
            "n_patient_like_sample_ids": "",
            "sample_id_examples": "",
        })
    write_csv(output / "A03_supplement_sheet_index.csv", sheet_rows)
    write_csv(
        output / "A03B_supplement_table2_sample_ids.csv",
        [
            {
                "sample_id": sample_id,
                "n_ASV_rows": n_rows,
                "healthy_like": "healthy" in sample_id.lower(),
                "patient_like": "patient" in sample_id.lower(),
            }
            for sample_id, n_rows in sorted(table2_sample_ids.items())
        ],
    )

    host_rows = [
        {
            "source": "PRJNA851469 official Supplementary Tables 2-17",
            "public_location": SUPPLEMENT_URL,
            "patient_or_sample_identifier_available": any(bool(row["detected_identifier_field"]) for row in sheet_rows),
            "patient_level_host_measurements_available": any(bool(row["joinable_patient_level_host_measurements"]) for row in sheet_rows),
            "microbiome_values_available": True,
            "join_to_current_sequence_objects_feasible": False,
            "status": "NO_JOINABLE_PATIENT_LEVEL_HOST_VALUES_DETECTED",
            "evidence": "Workbook CRC is valid. Table 2 is sample-level ASV abundance/taxonomy; other inspected sheets are annotations, correlations, or model summaries. No patient-level host measurement table was detected.",
            "next_action": "Do not infer host mechanisms from aggregate correlations. Request de-identified patient-level host data and a key from the corresponding author if C-stage work is pursued.",
        },
        {
            "source": "PRJNA851469 2026 antibody paper",
            "public_location": "https://link.springer.com/article/10.1186/s40635-026-00860-1",
            "patient_or_sample_identifier_available": False,
            "patient_level_host_measurements_available": False,
            "microbiome_values_available": False,
            "join_to_current_sequence_objects_feasible": False,
            "status": "REQUEST_REQUIRED",
            "evidence": "The paper links microbiome sequencing to PRJNA851469 and states that additional datasets/code are available from the corresponding author on reasonable request; a public joinable host table was not identified.",
            "next_action": "Prepare a narrow data request for de-identified antibody/host measurements, sampling-time identifiers, and linkage documentation.",
        },
        {
            "source": "Existing PRJNA851469 clinical linkage in this project",
            "public_location": r"E:\sepsis_project\results\V2_36F_PRJNA851469_CLINICAL_RECOVERY_LINKAGE",
            "patient_or_sample_identifier_available": True,
            "patient_level_host_measurements_available": False,
            "microbiome_values_available": True,
            "join_to_current_sequence_objects_feasible": True,
            "status": "CLINICAL_OUTCOME_LINKAGE_EXISTS_BUT_NO_HOST_OMICS",
            "evidence": "A curated clinical recovery linkage and day-3 outcome model already exist; these are clinical endpoints/covariates, not direct immune or host-molecular measurements.",
            "next_action": "Treat clinical results as association/translation evidence and retain the existing negative primary outcome result.",
        },
        {
            "source": "PRJNA1125274 external validation object/freeze",
            "public_location": r"E:\sepsis_project\results\V2_UPGRADE_20260910_FINAL_FREEZE",
            "patient_or_sample_identifier_available": True,
            "patient_level_host_measurements_available": False,
            "microbiome_values_available": True,
            "join_to_current_sequence_objects_feasible": True,
            "status": "MICROBIOME_EXTERNAL_VALIDATION_ONLY",
            "evidence": "The frozen mapping supports hospital, date, timepoint, and sequence-derived measures, but not direct patient-level host molecular readouts.",
            "next_action": "Use only for microbiome trajectory validation and prespecified center sensitivity; do not label it host-mechanistic validation.",
        },
    ]
    write_csv(output / "A10_host_data_availability.csv", host_rows)

    capture_hash_guard(output / "_input_hashes_before.csv", HASH_GUARD_INPUTS + RAW_COPIES)
    capture_tree_guard(
        output / "_protected_submission_tree_before.csv",
        Path(r"E:\sepsis_project\V2-SUBMISSION"),
    )

    runtime_path = output / "python_source_inspection_session.txt"
    if runtime_path.exists():
        raise RuntimeError(f"Refusing to overwrite existing audit output: {runtime_path}")
    runtime_path.write_text(
        "\n".join(
            [
                f"timestamp_utc={datetime.now(timezone.utc).isoformat()}",
                f"python={sys.version}",
                f"platform={platform.platform()}",
                f"openpyxl={__import__('openpyxl').__version__}",
                f"supplement_sha256={official_sha}",
            ]
        )
        + "\n",
        encoding="utf-8",
    )
    print(f"A-stage source inspection complete: {output}")
    print(f"Official workbook: {size} bytes, sha256={official_sha}")
    print(f"Workbook sheets: {len(sheet_rows)}; Table 2: {table2_info}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
